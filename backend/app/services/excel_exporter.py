"""
Сервис для экспорта данных в Excel
"""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session
from typing import Optional
from app.models.process_record import ProcessRecord


def export_to_excel(
    db: Session,
    table_source: Optional[str] = None
) -> Workbook:
    """
    Экспортирует данные в Excel файл.
    
    Args:
        db: сессия БД
        table_source: опциональный фильтр по источнику
    
    Returns:
        Workbook объект openpyxl
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Process Records"
    
    # Заголовки
    headers = [
        "ID",
        "Наименование процесса",
        "Продукт",
        "Статус",
        "Дата КБ",
        "ФИО клиента",
        "Дата приемки Банк",
        "ФИО от Банка",
        "НПП",
        "Сотрудник банка",
        "Комментарий",
        "Источник",
        "Дата импорта",
        "Создано",
        "Обновлено"
    ]
    
    # Стили для заголовков
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    # Записываем заголовки
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Получаем данные
    query = db.query(ProcessRecord).filter(ProcessRecord.is_archived == False)
    if table_source:
        query = query.filter(ProcessRecord.table_source == table_source)
    
    records = query.order_by(ProcessRecord.created_at.desc()).all()
    
    # Записываем данные
    for row_idx, record in enumerate(records, start=2):
        ws.cell(row=row_idx, column=1, value=str(record.id))
        ws.cell(row=row_idx, column=2, value=record.process_name)
        ws.cell(row=row_idx, column=3, value=record.product_type)
        ws.cell(row=row_idx, column=4, value=record.status)
        ws.cell(row=row_idx, column=5, value=record.date_kb.isoformat() if record.date_kb else None)
        ws.cell(row=row_idx, column=6, value=record.fio_customer)
        ws.cell(row=row_idx, column=7, value=record.date_bank_receipt.isoformat() if record.date_bank_receipt else None)
        ws.cell(row=row_idx, column=8, value=record.fio_bank_officer)
        ws.cell(row=row_idx, column=9, value=record.process_number)
        ws.cell(row=row_idx, column=10, value=record.bank_employee_name)
        ws.cell(row=row_idx, column=11, value=record.comments)
        ws.cell(row=row_idx, column=12, value=record.table_source)
        ws.cell(row=row_idx, column=13, value=record.import_date.isoformat() if record.import_date else None)
        ws.cell(row=row_idx, column=14, value=record.created_at.isoformat() if record.created_at else None)
        ws.cell(row=row_idx, column=15, value=record.updated_at.isoformat() if record.updated_at else None)
    
    # Настройка ширины колонок
    column_widths = [36, 30, 20, 15, 12, 25, 12, 25, 8, 25, 40, 20, 20, 20, 20]
    for col_idx, width in enumerate(column_widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    
    # Автофильтр
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(records) + 1}"
    
    return wb

