#!/usr/bin/env python3
"""Скрипт для анализа структуры Excel файлов"""
import openpyxl
import sys
import os

excel_files = [
    "!Чек-лист Аренда сейфовых ячеек.xlsx",
    "!Чек-лист Депозиты ФЛ.xlsx",
    "!Чек-Лист ДЮЛ.xlsx",
    "!Чек-Лист Кредиты.xlsx",
    "!Чек-Лист РКО.xlsx",
    "!Чек-Лист РЦ.xlsx",
]

base_path = "/Users/danil/My Drive/IDF Eurasia/Чек-листы"

for filename in excel_files:
    filepath = os.path.join(base_path, filename)
    if not os.path.exists(filepath):
        print(f"⚠️  Файл не найден: {filename}")
        continue
    
    try:
        wb = openpyxl.load_workbook(filepath)
        ws = wb.active
        
        headers = [cell.value for cell in ws[1] if cell.value]
        
        print(f"\n📄 {filename}")
        print(f"   Колонки ({len(headers)}):")
        for i, h in enumerate(headers, 1):
            print(f"   {i}. '{h}'")
        
        # Показываем первую строку данных
        if ws.max_row > 1:
            print(f"   Первая строка данных:")
            row2 = [cell.value for cell in ws[2]]
            for i, (header, value) in enumerate(zip(headers, row2[:len(headers)])):
                if value:
                    print(f"   {header}: {value}")
        
    except Exception as e:
        print(f"❌ Ошибка при чтении {filename}: {e}")

